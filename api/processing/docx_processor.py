"""DOCX document processing for text extraction."""

import logging
from pathlib import Path

from docx import Document
from docx.table import Table as DocxTable

logger = logging.getLogger(__name__)


class DocxProcessor:
    """Process DOCX files and extract text as markdown."""

    def extract(self, file_path: Path) -> str:
        """Extract text from a DOCX file and convert to markdown.

        Args:
            file_path: Path to the DOCX file

        Returns:
            Markdown-formatted text content
        """
        logger.info(f"Extracting text from DOCX: {file_path}")

        doc = Document(str(file_path))
        markdown_parts = []
        markdown_parts.append(f"# {file_path.name}\n")

        for element in doc.element.body:
            # Handle paragraphs
            if element.tag.endswith("p"):
                para = self._find_paragraph(doc, element)
                if para:
                    formatted = self._format_paragraph(para)
                    if formatted:
                        markdown_parts.append(formatted)

            # Handle tables
            elif element.tag.endswith("tbl"):
                table = self._find_table(doc, element)
                if table:
                    formatted = self._format_table(table)
                    if formatted:
                        markdown_parts.append(formatted)

        markdown = "\n\n".join(markdown_parts)
        logger.info(f"Extracted {len(markdown)} characters from DOCX")
        return markdown

    def _find_paragraph(self, doc: Document, element) -> any:
        """Find the paragraph object for an element."""
        for para in doc.paragraphs:
            if para._element is element:
                return para
        return None

    def _find_table(self, doc: Document, element) -> DocxTable | None:
        """Find the table object for an element."""
        for table in doc.tables:
            if table._element is element:
                return table
        return None

    def _format_paragraph(self, para) -> str | None:
        """Format a paragraph as markdown.

        Args:
            para: Paragraph object from python-docx

        Returns:
            Markdown-formatted string
        """
        text = para.text.strip()
        if not text:
            return None

        style_name = para.style.name.lower() if para.style else ""

        # Detect headings
        if "heading 1" in style_name or "title" in style_name:
            return f"# {text}"
        elif "heading 2" in style_name:
            return f"## {text}"
        elif "heading 3" in style_name:
            return f"### {text}"
        elif "heading 4" in style_name:
            return f"#### {text}"
        elif "heading 5" in style_name:
            return f"##### {text}"
        elif "heading 6" in style_name:
            return f"###### {text}"

        # Detect list items
        elif "list" in style_name or "bullet" in style_name:
            return f"- {text}"
        elif "number" in style_name:
            return f"1. {text}"

        # Regular paragraph
        return text

    def _format_table(self, table: DocxTable) -> str | None:
        """Format a table as markdown.

        Args:
            table: Table object from python-docx

        Returns:
            Markdown-formatted table string
        """
        if not table.rows:
            return None

        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("|", "\\|") for cell in row.cells]
            rows.append("| " + " | ".join(cells) + " |")

        if len(rows) < 1:
            return None

        # Add separator after header row
        num_cols = len(table.rows[0].cells)
        separator = "| " + " | ".join(["---"] * num_cols) + " |"

        result = [rows[0], separator] + rows[1:]
        return "\n".join(result)


# Global instance
docx_processor = DocxProcessor()
