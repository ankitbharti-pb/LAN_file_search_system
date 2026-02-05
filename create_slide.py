"""Generate a one-page PowerPoint slide for the LAN File Search System — use case focused."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

# --- Constants ---
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
ACCENT_BLUE = RGBColor(0x2A, 0x5E, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MEDIUM_GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_BG = RGBColor(0xF5, 0xF7, 0xFA)
QUERY_BG = RGBColor(0xEE, 0xF2, 0xF7)
BORDER_BLUE = RGBColor(0xC8, 0xD8, 0xE8)
SOURCE_GREEN = RGBColor(0x1A, 0x7A, 0x4C)

ROOT = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(ROOT, "LAN_File_Search_System_v2.pptx")


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_tb(slide, left, top, width, height, word_wrap=True):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tb.text_frame.word_wrap = word_wrap
    return tb


def styled_run(paragraph, text, size, color, bold=False, italic=False, font="Calibri"):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    return run


def add_para(tf, text, size, color, bold=False, space_after=Pt(3), space_before=Pt(0), alignment=PP_ALIGN.LEFT):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.space_after = space_after
    p.space_before = space_before
    p.alignment = alignment
    p.font.name = "Calibri"
    return p


def set_first_para(tf, text, size, color, bold=False, space_after=Pt(3), alignment=PP_ALIGN.LEFT):
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.space_after = space_after
    p.alignment = alignment
    p.font.name = "Calibri"
    return p


def build_use_case(slide, y_top, box_h, label, query_text, response_lines, source_text):
    """Build one use-case row: label above, query box on left, response box on right."""
    margin = Inches(0.6)
    gap = Inches(0.3)
    query_w = Inches(3.8)
    response_w = SLIDE_WIDTH - margin - margin - query_w - gap
    inner_pad = Inches(0.15)

    # ── Section label ──
    tb_label = add_tb(slide, margin, y_top, Inches(12), Inches(0.3))
    set_first_para(tb_label.text_frame, label, 12, ACCENT_BLUE, bold=True, space_after=Pt(0))

    row_top = y_top + Inches(0.32)

    # ── Query box (left) ──
    add_rect(slide, margin, row_top, query_w, box_h, QUERY_BG, BORDER_BLUE)

    tb_q = add_tb(slide, margin + inner_pad, row_top + inner_pad, query_w - 2 * inner_pad, box_h - 2 * inner_pad)
    tf_q = tb_q.text_frame
    tf_q.vertical_anchor = MSO_ANCHOR.MIDDLE

    set_first_para(tf_q, "USER QUERY", 9, MEDIUM_GRAY, bold=True, space_after=Pt(8))
    # Query text in quotes, larger
    p_query = tf_q.add_paragraph()
    styled_run(p_query, "\u201c", 18, ACCENT_BLUE, bold=True)
    styled_run(p_query, query_text, 13, DARK_BLUE, bold=False, italic=True)
    styled_run(p_query, "\u201d", 18, ACCENT_BLUE, bold=True)
    p_query.space_after = Pt(0)
    p_query.space_before = Pt(4)

    # ── Response box (right) ──
    resp_left = margin + query_w + gap
    add_rect(slide, resp_left, row_top, response_w, box_h, WHITE, BORDER_BLUE)

    # Blue left accent bar inside response box
    add_rect(slide, resp_left, row_top, Pt(4), box_h, ACCENT_BLUE)

    tb_r = add_tb(slide, resp_left + Inches(0.2), row_top + inner_pad,
                  response_w - Inches(0.35), box_h - 2 * inner_pad)
    tf_r = tb_r.text_frame
    tf_r.vertical_anchor = MSO_ANCHOR.TOP

    set_first_para(tf_r, "AI ANSWER", 9, MEDIUM_GRAY, bold=True, space_after=Pt(6))

    for line in response_lines:
        if line.startswith("##"):
            # Sub-heading
            add_para(tf_r, line[2:].strip(), 10, DARK_BLUE, bold=True, space_after=Pt(2), space_before=Pt(5))
        elif line.startswith("\u2022"):
            add_para(tf_r, line, 9.5, DARK_GRAY, space_after=Pt(2), space_before=Pt(0))
        else:
            add_para(tf_r, line, 9.5, DARK_GRAY, space_after=Pt(3))

    # Source line
    p_src = tf_r.add_paragraph()
    styled_run(p_src, "Source: ", 9, MEDIUM_GRAY, bold=True)
    styled_run(p_src, source_text, 9, SOURCE_GREEN, bold=False)
    p_src.space_before = Pt(6)
    p_src.space_after = Pt(0)


def build_slide():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ── Header banner ──────────────────────────────────────────────
    banner_h = Inches(1.05)
    add_rect(slide, Inches(0), Inches(0), SLIDE_WIDTH, banner_h, DARK_BLUE)
    add_rect(slide, Inches(0), banner_h, SLIDE_WIDTH, Pt(3), ACCENT_BLUE)

    tb_title = add_tb(slide, Inches(0.6), Inches(0.15), Inches(10), Inches(0.5))
    set_first_para(tb_title.text_frame, "LAN FILE SEARCH SYSTEM", 26, WHITE, bold=True, space_after=Pt(0))

    tb_sub = add_tb(slide, Inches(0.6), Inches(0.62), Inches(10), Inches(0.35))
    set_first_para(tb_sub.text_frame,
                   "AI-Powered Enterprise Document Search  |  Runs Entirely on LAN",
                   13, RGBColor(0xBB, 0xCC, 0xDD), space_after=Pt(0))

    # ── Use Case 1: Policy Document Search ─────────────────────────
    uc1_response = [
        "The hospitalization benefit under ReAssure 2.0 covers expenses related to reaching a hospital, "
        "during hospitalization, and before and after hospitalization, up to the Base Sum Insured for a "
        "single claim (Clause: Note, ReAssure-2.0-Policy-Wording.pdf). Additionally, the policy includes "
        "coverage for expenses incurred 60 days prior to admission and 180 days after discharge if related "
        "to the same condition (Clause: 4.2.2, ReAssure-2.0-Policy-Wording.pdf).",
        "## Charges Covered",
        "\u2022  Expenses in reaching the hospital",
        "\u2022  Hospitalization expenses",
        "\u2022  Pre- and post-hospitalization expenses",
        "\u2022  Home care/domiciliary treatment",
        "\u2022  Organ donor expenses",
        "## Additional Benefits",
        "\u2022  Hospital Cash: A fixed amount per day for hospitalization, up to 30 days, payable if "
        "hospitalized for 48 hours or more (Clause: 4.14, ReAssure-2.0-Policy-Wording.pdf).",
        "\u2022  Shared Accommodation Cash Benefit: Additional daily amount if opting for a shared room "
        "(Clause: 4.11, ReAssure-2.0-Policy-Wording.pdf).",
        "## Important Notes",
        "\u2022  Expenses in reaching hospital and pre/post-hospitalization for the first-ever hospitalization "
        "are considered as part of the initial claim (Note: Year 1).",
        "\u2022  The maximum payable for any claim is the Base Sum Insured (Note: Year 1).",
        "\u2022  Charges for ICU, organ donor, and other specific expenses are also covered as per the detailed "
        "policy provisions, but the core hospitalization benefit primarily includes the above components.",
    ]

    build_use_case(
        slide,
        y_top=Inches(1.3),
        box_h=Inches(3.85),
        label="USE CASE 1  \u2014  Policy Document Search",
        query_text="hospitalization benefit and charges for reassure2.0",
        response_lines=uc1_response,
        source_text="ReAssure-2.0-Policy-Wording.pdf (10 chunks)",
    )

    # ── Use Case 2: SOP Document Search ────────────────────────────
    uc2_response = [
        "To change your mobile number with Niva, you need to follow the process via PB MyAccount, "
        "where you can complete the update through OTP verification.",
        "The change is processed as an instant endorsement, and an endorsement letter (EL) is sent "
        "to your registered email.",
    ]

    build_use_case(
        slide,
        y_top=Inches(5.4),
        box_h=Inches(1.75),
        label="USE CASE 2  \u2014  SOP Document Search",
        query_text="mobile number change sop for niva",
        response_lines=uc2_response,
        source_text="updated_sop2.csv (10 chunks)",
    )

    # ── Footer bar ─────────────────────────────────────────────────
    footer_h = Inches(0.35)
    footer_top = SLIDE_HEIGHT - footer_h
    add_rect(slide, Inches(0), footer_top, SLIDE_WIDTH, footer_h, DARK_BLUE)
    tb_footer = add_tb(slide, Inches(0.6), footer_top + Inches(0.05), Inches(12), Inches(0.25))
    set_first_para(tb_footer.text_frame,
                   "FastAPI + React  |  FAISS + BM25 Hybrid Search  |  Sentence Transformers  |  OpenAI / Gemini",
                   9, RGBColor(0xAA, 0xBB, 0xCC), space_after=Pt(0), alignment=PP_ALIGN.CENTER)

    prs.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    build_slide()
